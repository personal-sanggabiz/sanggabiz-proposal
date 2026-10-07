"""Helper library: gaya compro Sanggabiz (navy, kuning, teal, Montserrat) untuk python-pptx."""
import os, copy
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")
ICONS = os.path.join(ASSETS, "icons")
os.makedirs(ICONS, exist_ok=True)

# ---- design tokens (dari analisis compro) ----
NAVY_L = "0B396D"
NAVY_D = "092647"
NAVY_P = "0B3361"   # panel
NAVY_DD = "0D2A4E"
YELLOW = "FFC00C"
TEAL = "2F7E9D"
WHITE = "FFFFFF"
GREY = "B0C0CF"
BLUE = "2446BC"
INK = "1D1D1F"

F_XB = "Montserrat ExtraBold"
F_SB = "Montserrat SemiBold"
F_M = "Montserrat Medium"
F_R = "Montserrat"

SW, SH = 13.333, 7.5
ML = 0.65
CW = SW - 2 * ML


def rgb(h):
    return RGBColor.from_string(h)


def new_prs():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    return prs


def blank(prs, bg="grad"):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    if bg == "grad":
        f = s.background.fill
        f.gradient()
        f.gradient_angle = 300
        st = f.gradient_stops
        st[0].position = 0.0
        st[0].color.rgb = rgb(NAVY_L)
        st[1].position = 1.0
        st[1].color.rgb = rgb(NAVY_D)
    elif bg == "white":
        f = s.background.fill
        f.solid()
        f.fore_color.rgb = rgb(WHITE)
    return s


def _set_alpha(shape_fill_elm, alpha):
    clr = shape_fill_elm.find(".//" + qn("a:srgbClr"))
    if clr is not None:
        a = etree.SubElement(clr, qn("a:alpha"))
        a.set("val", str(int(alpha * 1000)))


def rect(s, x, y, w, h, fill=None, line=None, lw=1.5, radius=None, alpha=None, shape=None, name=None):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE)
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, radius / min(w, h)) if radius <= 0.5 * min(w, h) else 0.5
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            _set_alpha(sh._element.spPr.find(qn("a:solidFill")), alpha)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    if name:
        sh.name = name
    # remove text placeholder style influence
    sh.text_frame.text = ""
    return sh


def pill(s, x, y, w, h, fill=YELLOW, line=None, lw=1.5, alpha=None):
    return rect(s, x, y, w, h, fill=fill, line=line, lw=lw, radius=h / 2, alpha=alpha)


def _apply_run(r, size, bold, color, font, italic=False):
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = rgb(color)
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", font)


def text(s, x, y, w, h, paras, size=14, bold=False, color=WHITE, font=F_R, align="l", anchor="t",
         spacing=1.1, after=4, bullet=None, margin=0.0, italic=False, caps=False, tracking=None, shape=None):
    """paras: str | list[str | list[(text, opts)]]. opts: bold,color,size,font,italic."""
    if shape is None:
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        tb = shape
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = [paras]
    first = True
    for p_ in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT, "j": PP_ALIGN.JUSTIFY}[align]
        p.line_spacing = spacing
        p.space_after = Pt(after)
        runs = [(p_, {})] if isinstance(p_, str) else p_
        is_bullet = bullet
        for t, o in runs:
            if isinstance(o, str):
                o = {}
            if o.get("nobullet"):
                is_bullet = None
            r = p.add_run()
            r.text = t.upper() if caps else t
            _apply_run(r, o.get("size", size), o.get("bold", bold), o.get("color", color),
                       o.get("font", font), o.get("italic", italic))
            if tracking:
                r._r.get_or_add_rPr().set("spc", str(tracking))
        if is_bullet:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches(0.22))))
            pPr.set("indent", str(int(-Inches(0.22))))
            bc = etree.SubElement(pPr, qn("a:buClr"))
            c = etree.SubElement(bc, qn("a:srgbClr"))
            c.set("val", YELLOW)
            bf = etree.SubElement(pPr, qn("a:buFont"))
            bf.set("typeface", "Arial")
            bu = etree.SubElement(pPr, qn("a:buChar"))
            bu.set("char", is_bullet if isinstance(is_bullet, str) else "•")
    return tb


def photo(s, path, x, y, w, h, rounded=True, rot=0, focus=(0.5, 0.5), line=None):
    """Sisip foto, dipotong ke rasio w:h (fokus 0..1)."""
    im = Image.open(path).convert("RGB")
    tw, th = w / h, None
    iw, ih = im.size
    if iw / ih > w / h:
        nw = int(ih * w / h)
        left = int((iw - nw) * focus[0])
        im = im.crop((left, 0, left + nw, ih))
    else:
        nh = int(iw * h / w)
        top = int((ih - nh) * focus[1])
        im = im.crop((0, top, iw, top + nh))
    im.thumbnail((1600, 1600))
    out = os.path.join(ASSETS, "cache_" + os.path.basename(path).rsplit(".", 1)[0] + f"_{int(w*100)}x{int(h*100)}.jpg")
    im.save(out, quality=85)
    pic = s.shapes.add_picture(out, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
    if rot:
        pic.rotation = rot
    if line:
        pic.line.color.rgb = rgb(line)
        pic.line.width = Pt(2.5)
    return pic


def logo(s, x, y, h=0.32):
    p = os.path.join(ASSETS, "logo_white.png")
    im = Image.open(p)
    w = h * im.size[0] / im.size[1]
    return s.shapes.add_picture(p, Inches(x), Inches(y), Inches(w), Inches(h))


# ---------------- ikon garis putih ----------------
def _icon_canvas():
    S = 4
    im = Image.new("RGBA", (256 * S, 256 * S), (0, 0, 0, 0))
    return im, ImageDraw.Draw(im), S


def make_icon(name, color=WHITE):
    path = os.path.join(ICONS, f"{name}_{color}.png")
    if os.path.exists(path):
        return path
    im, d, S = _icon_canvas()
    c = tuple(int(color[i:i + 2], 16) for i in (0, 2, 4)) + (255,)
    lw = 13 * S

    def L(pts, w=lw):
        d.line([(x * S, y * S) for x, y in pts], fill=c, width=w, joint="curve")
        for x, y in (pts[0], pts[-1]):
            d.ellipse((x * S - w / 2, y * S - w / 2, x * S + w / 2, y * S + w / 2), fill=c)

    def RR(b, r=18, w=lw):
        d.rounded_rectangle([v * S for v in b], radius=r * S, outline=c, width=w)

    def EL(b, w=lw, fill=None):
        d.ellipse([v * S for v in b], outline=c, width=w, fill=fill)

    if name == "calendar":
        RR((36, 52, 220, 220)); L([(36, 100), (220, 100)]); L([(84, 30), (84, 70)]); L([(172, 30), (172, 70)])
        for i in range(3):
            for j in range(2):
                d.rectangle([(70 + i * 40) * S, (124 + j * 40) * S, (90 + i * 40) * S, (144 + j * 40) * S], fill=c)
    elif name == "people":
        EL((88, 40, 168, 120)); d.arc([50 * S, 130 * S, 206 * S, 290 * S], 180, 360, fill=c, width=lw)
        L([(50, 210), (206, 210)])
    elif name == "doc":
        RR((60, 30, 196, 226)); L([(92, 90), (164, 90)]); L([(92, 130), (164, 130)]); L([(92, 170), (140, 170)])
    elif name == "chart":
        L([(40, 216), (216, 216)]); L([(40, 40), (40, 216)])
        for i, h in enumerate((70, 110, 150)):
            d.rectangle([(70 + i * 46) * S, (216 - h) * S, (98 + i * 46) * S, 216 * S], fill=c)
    elif name == "pin":
        d.arc([60 * S, 30 * S, 196 * S, 166 * S], 150, 390, fill=c, width=lw)
        L([(72, 140), (128, 226), (184, 140)]); EL((104, 76, 152, 124))
    elif name == "check":
        EL((36, 36, 220, 220)); L([(84, 130), (116, 162), (176, 96)])
    elif name == "clock":
        EL((36, 36, 220, 220)); L([(128, 76), (128, 128), (168, 152)])
    elif name == "database":
        d.ellipse([56 * S, 30 * S, 200 * S, 80 * S], outline=c, width=lw)
        L([(56, 55), (56, 200)]); L([(200, 55), (200, 200)])
        d.arc([56 * S, 150 * S, 200 * S, 226 * S], 0, 180, fill=c, width=lw)
        d.arc([56 * S, 100 * S, 200 * S, 176 * S], 0, 180, fill=c, width=lw)
    elif name == "target":
        EL((32, 32, 224, 224)); EL((80, 80, 176, 176)); d.ellipse([116 * S, 116 * S, 140 * S, 140 * S], fill=c)
    elif name == "search":
        EL((40, 40, 172, 172)); L([(160, 160), (218, 218)])
    elif name == "route":
        EL((30, 150, 80, 200), fill=None); EL((176, 50, 226, 100))
        L([(80, 175), (140, 175), (140, 75), (176, 75)])
    elif name == "shield":
        L([(128, 30), (210, 62), (210, 130), (128, 226), (46, 130), (46, 62), (128, 30)]); L([(92, 128), (120, 156), (170, 100)])
    elif name == "layers":
        L([(128, 40), (224, 90), (128, 140), (32, 90), (128, 40)]); L([(32, 130), (128, 180), (224, 130)]); L([(32, 170), (128, 220), (224, 170)])
    elif name == "flag":
        L([(60, 30), (60, 226)]); L([(60, 40), (200, 40), (170, 84), (200, 128), (60, 128)])
    elif name == "link":
        RR((30, 90, 130, 166), r=38); RR((126, 90, 226, 166), r=38)
    elif name == "users":
        EL((40, 70, 110, 140)); EL((146, 70, 216, 140)); d.arc([20 * S, 140 * S, 130 * S, 270 * S], 180, 360, fill=c, width=lw)
        d.arc([126 * S, 140 * S, 236 * S, 270 * S], 180, 360, fill=c, width=lw)
    elif name == "lock":
        RR((50, 110, 206, 224), r=20); d.arc([76 * S, 30 * S, 180 * S, 170 * S], 180, 360, fill=c, width=lw)
        L([(76, 100), (76, 110)]); L([(180, 100), (180, 110)])
    elif name == "book":
        L([(128, 60), (128, 216)]); L([(128, 60), (60, 44), (30, 60), (30, 200), (60, 184), (128, 216)])
        L([(128, 60), (196, 44), (226, 60), (226, 200), (196, 184), (128, 216)])
    im = im.resize((256, 256), Image.LANCZOS)
    im.save(path)
    return path


def icon(s, name, x, y, size=0.5, color=WHITE):
    return s.shapes.add_picture(make_icon(name, color), Inches(x), Inches(y), Inches(size), Inches(size))


# ---------------- tabel ----------------
def _border(cell, color=TEAL, w=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = etree.SubElement(tcPr, qn(tag))
        ln.set("w", str(int(Pt(w))))
        sf = etree.SubElement(ln, qn("a:solidFill"))
        c = etree.SubElement(sf, qn("a:srgbClr"))
        c.set("val", color)
    # borders must come before fill in schema order -> move fill to end
    for tag in ("a:solidFill", "a:noFill"):
        f = tcPr.find(qn(tag))
        if f is not None:
            tcPr.remove(f)
            tcPr.append(f)


def table(s, x, y, colw, rows, size=12, head_size=12, row_h=0.4, head_h=0.42, first_bold=True,
          head_fill=YELLOW, head_color=NAVY_D, fills=(NAVY_P, NAVY_DD), aligns=None, valign="m", after=2):
    n_r, n_c = len(rows), len(colw)
    gs = s.shapes.add_table(n_r, n_c, Inches(x), Inches(y), Inches(sum(colw)), Inches(head_h + row_h * (n_r - 1)))
    tbl = gs.table
    tblPr = tbl._tbl.tblPr
    for a in ("firstRow", "bandRow"):
        tblPr.set(a, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # no style, no grid
    for i, w in enumerate(colw):
        tbl.columns[i].width = Inches(w)
    for r in range(n_r):
        tbl.rows[r].height = Inches(head_h if r == 0 else row_h)
        for c in range(n_c):
            cell = tbl.cell(r, c)
            val = rows[r][c]
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = {"m": MSO_ANCHOR.MIDDLE, "t": MSO_ANCHOR.TOP}[valign]
            tf = cell.text_frame
            tf.word_wrap = True
            paras = val if isinstance(val, list) else [val]
            for k, pv in enumerate(paras):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                p.space_after = Pt(after)
                p.line_spacing = 1.05
                al = (aligns[c] if aligns else "l")
                p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[al]
                runs = [(pv, {})] if isinstance(pv, str) else ([pv] if isinstance(pv, tuple) else pv)
                for t, o in runs:
                    rr = p.add_run()
                    rr.text = t
                    if r == 0:
                        _apply_run(rr, head_size, True, head_color, F_R)
                    else:
                        _apply_run(rr, o.get("size", size), o.get("bold", first_bold and c == 0),
                                   o.get("color", WHITE), o.get("font", F_R), o.get("italic", False))
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(head_fill if r == 0 else fills[(r - 1) % 2])
            _border(cell)
    return gs


# ---------------- komponen slide ----------------
def header(s, kicker, title, tsize=28, w=None):
    w = w or CW
    if kicker:
        text(s, ML, 0.42, w, 0.3, kicker, size=11.5, bold=True, color=YELLOW, font=F_R, caps=True, tracking=120)
    text(s, ML, 0.72, w, 1.0, title, size=tsize, color=WHITE, font=F_XB, spacing=1.0, after=0)


def footer(s, n, total=None):
    logo(s, ML, 7.02, 0.26)
    text(s, ML + 1.45, 7.06, 7.5, 0.25, "Proposal Teknis  ·  RFP No. 0648/TF-JKT/VIII/2026  ·  Tanoto Foundation",
         size=9, color=GREY, font=F_R)
    text(s, SW - ML - 1.0, 7.06, 1.0, 0.25, str(n), size=10, bold=True, color=YELLOW, align="r")


def tab(s, x, y, w, label, fill=YELLOW, color=NAVY_D, h=0.4, size=13):
    p = pill(s, x, y, w, h, fill=fill)
    text(s, x, y, w, h, label, size=size, bold=True, color=color, font=F_R, align="c", anchor="m", shape=None)
    return p


def card(s, x, y, w, h, fill=TEAL, line=WHITE, lw=2.25, radius=0.28, alpha=None):
    return rect(s, x, y, w, h, fill=fill, line=line, lw=lw, radius=radius, alpha=alpha)
