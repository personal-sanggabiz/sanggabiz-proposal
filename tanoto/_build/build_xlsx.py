"""Model anggaran XLSX (formula hidup) dari asumsi yang sama dengan deck."""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import budget as B
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

OUT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "Model Anggaran - Tanoto PPG LPTK.xlsx"))
wb = Workbook()
YEL = "FFC00C"
hf = Font(bold=True, color="092647")
hfill = PatternFill("solid", fgColor=YEL)
inp = PatternFill("solid", fgColor="FFF4C2")
thin = Side(style="thin", color="B0C0CF")
bd = Border(left=thin, right=thin, top=thin, bottom=thin)
rp = '"Rp"#,##0'

ws = wb.active
ws.title = "Asumsi"
ws["A1"] = "ASUMSI (sel kuning dapat diubah; harga satuan adalah ASUMSI yang perlu diverifikasi)"
ws["A1"].font = Font(bold=True, size=12)
rows = [
    ("Pagu maksimum", B.PAGU, "RFP bagian E"),
    ("Personel per kunjungan", B.PERSONS, "2 TPN + 1 mitra (RFP)"),
    ("Tarif akomodasi per orang per malam", B.RATE_NIGHT, "RFP"),
    ("Malam per kunjungan standar", B.NIGHTS_STD, "RFP"),
    ("Malam tambahan per LPTK berikutnya dalam klaster", B.CLUSTER_EXTRA_NIGHTS, "Rancangan klaster"),
    ("Tiket PP rata-rata per orang", B.TICKET_RT, "ASUMSI: verifikasi rute dan harga aktual"),
    ("Transportasi lokal per hari", B.LOCAL_DAY, "ASUMSI: mobil, driver, BBM"),
    ("Hari transportasi lokal per kunjungan", B.LOCAL_DAYS, "ASUMSI"),
    ("Uang harian dan logistik per orang per hari", B.DAILY, "ASUMSI"),
    ("Honorarium TPN per orang per kunjungan", B.HONOR_TPN, "ASUMSI: acuan SBM Kementerian, tanyakan ke TF"),
    ("Jumlah TPN per kunjungan", B.N_TPN, "RFP"),
    ("Platform, repositori, alat data (paket)", B.TOOLS, "ASUMSI"),
    ("Kontingensi perjalanan", 9_330_000, "Penyeimbang agar total sama dengan pagu"),
    ("Management fee (bersih)", 42_000_000, "Bruto 8-10% dari total (band RFP)"),
    ("Tarif PPN atas fee", B.PPN, "ASUMSI: konfirmasi jenis pajak ke TF"),
    ("Jumlah LPTK inti", 12, "RFP"),
    ("Jumlah LPTK tambahan", 2, "Opsional (RFP membolehkan)"),
    ("Jumlah klaster rute (perjalanan PP)", 5, "Medan, Pekanbaru, Jambi, Semarang-Salatiga-Kudus, Samarinda"),
    ("Total malam LPTK inti (klaster)", sum(B.NIGHTS_STD + B.CLUSTER_EXTRA_NIGHTS * (k - 1) for _, k in B.CLUSTERS), "9+5+3+9+3"),
]
for i, (l, v, n) in enumerate(rows, start=3):
    ws.cell(i, 1, l)
    c = ws.cell(i, 2, v)
    c.fill = inp
    c.number_format = "0.00" if isinstance(v, float) else "#,##0"
    ws.cell(i, 3, n)
ws.column_dimensions["A"].width = 52
ws.column_dimensions["B"].width = 18
ws.column_dimensions["C"].width = 60
g = {l: f"Asumsi!$B${i}" for i, (l, v, n) in enumerate(rows, start=3)}

w2 = wb.create_sheet("RAB Usulan (14 LPTK)")
w2["A1"] = "RENCANA ANGGARAN BIAYA (usulan utama: 12 LPTK inti + LPTK tambahan)"
w2["A1"].font = Font(bold=True, size=12)
for j, h in enumerate(["No", "Komponen", "Satuan biaya", "Volume", "Durasi/Pengali", "Subtotal"], 1):
    c = w2.cell(3, j, h)
    c.font = hf
    c.fill = hfill
    c.border = bd
inti, tamb = g["Jumlah LPTK inti"], g["Jumlah LPTK tambahan"]
lines = [
    ("Honorarium TPN", f"={g['Honorarium TPN per orang per kunjungan']}", f"={g['Jumlah TPN per kunjungan']}", f"={inti}+{tamb}"),
    ("Akomodasi (malam-orang)", f"={g['Tarif akomodasi per orang per malam']}", f"={g['Personel per kunjungan']}",
     f"={g['Total malam LPTK inti (klaster)']}+{tamb}*{g['Malam per kunjungan standar']}"),
    ("Tiket pulang-pergi", f"={g['Tiket PP rata-rata per orang']}", f"={g['Personel per kunjungan']}",
     f"={g['Jumlah klaster rute (perjalanan PP)']}+{tamb}"),
    ("Transportasi lokal (hari kunjungan)", f"={g['Transportasi lokal per hari']}", f"={inti}+{tamb}", f"={g['Hari transportasi lokal per kunjungan']}"),
    ("Uang harian dan logistik", f"={g['Uang harian dan logistik per orang per hari']}",
     f"=({inti}+{tamb})*{g['Personel per kunjungan']}", f"={g['Hari transportasi lokal per kunjungan']}"),
    ("Platform, repositori, alat data", f"={g['Platform, repositori, alat data (paket)']}", 1, 1),
    ("Kontingensi perjalanan", f"={g['Kontingensi perjalanan']}", 1, 1),
]
r = 4
for n, (nm, price, vol, mult) in enumerate(lines, 1):
    w2.cell(r, 1, n)
    w2.cell(r, 2, nm)
    w2.cell(r, 3, price)
    w2.cell(r, 4, vol)
    w2.cell(r, 5, mult)
    w2.cell(r, 6, f"=C{r}*D{r}*E{r}")
    r += 1
last = r - 1
w2.cell(r, 2, "Subtotal perjalanan dan logistik").font = Font(bold=True)
w2.cell(r, 6, f"=SUM(F4:F{last})").font = Font(bold=True)
sub = r; r += 1
w2.cell(r, 2, "Management fee (termasuk honorarium mitra)")
w2.cell(r, 6, f"={g['Management fee (bersih)']}")
fee = r; r += 1
w2.cell(r, 2, "PPN atas management fee")
w2.cell(r, 6, f"=ROUND(F{fee}*{g['Tarif PPN atas fee']},0)")
ppn = r; r += 1
w2.cell(r, 2, "TOTAL").font = Font(bold=True)
w2.cell(r, 6, f"=F{sub}+F{fee}+F{ppn}").font = Font(bold=True)
tot = r; r += 1
w2.cell(r, 2, "Selisih terhadap pagu (harus nol atau positif)")
w2.cell(r, 6, f"={g['Pagu maksimum']}-F{tot}")
r += 1
w2.cell(r, 2, "Porsi perjalanan dan logistik (band RFP 90-92%)")
w2.cell(r, 6, f"=F{sub}/F{tot}")
pct1 = r; r += 1
w2.cell(r, 2, "Porsi fee + pajak (band RFP 8-10%)")
w2.cell(r, 6, f"=(F{fee}+F{ppn})/F{tot}")
pct2 = r
for row in w2.iter_rows(min_row=4, max_row=pct2):
    for c in row:
        c.border = bd
        if c.column in (3, 6):
            c.number_format = rp
w2.cell(pct1, 6).number_format = "0.00%"
w2.cell(pct2, 6).number_format = "0.00%"
for j, wd in enumerate([5, 50, 16, 12, 16, 20], 1):
    w2.column_dimensions[get_column_letter(j)].width = wd
for c in w2[tot]:
    c.fill = hfill

w3 = wb.create_sheet("Skenario 12 LPTK")
w3["A1"] = "Alternatif: hanya 12 LPTK inti (nilai tetap dari model; hitung ulang lewat sheet Asumsi bila perlu)"
w3["A1"].font = Font(bold=True)
b = B.B
data = [("Honorarium TPN", b["lines"]["honor"]), ("Akomodasi", b["lines"]["akom"]), ("Tiket", b["lines"]["tiket"]),
        ("Transportasi lokal", b["lines"]["lokal"]), ("Uang harian dan logistik", b["lines"]["harian"]),
        ("Platform dan alat data", b["lines"]["tools"]), ("Kontingensi", b["lines"]["kontingensi"]),
        ("Management fee", b["fee"]), ("PPN 11% atas fee", b["ppn"]), ("TOTAL", b["total"])]
for i, (l, v) in enumerate(data, start=3):
    w3.cell(i, 1, l)
    c = w3.cell(i, 2, v)
    c.number_format = rp
w3.column_dimensions["A"].width = 30
w3.column_dimensions["B"].width = 20
w3["A14"] = "Termin pembayaran (usulan): 30% Laporan Awal disetujui; 40% laporan pendampingan 12 LPTK; 30% Paket Laporan Akhir final."
wb.save(OUT)
print("saved", OUT)
